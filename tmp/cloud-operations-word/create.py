import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path.cwd(); out=ROOT/'output/docx/Rajendra_P_N_Infrastructure_Platform_Cloud_Operations.docx'
base=json.loads(Path('plugins/resume-creator-plugin/assets/people/rajendra-prasad-n.json').read_text())
profile={k:base[k] for k in ['person_id','full_name','headline','location','email','education','certifications']}
profile['summary']='Lead Azure DevOps Engineer with 10+ years of experience in enterprise infrastructure, cloud and platform operations, Kubernetes, Linux, and production support. Engineers Azure platforms with Terraform, Infrastructure as Code (IaC), Helm, and GitOps. Focuses on high availability, scalability, capacity management, observability, security hardening, and backup and recovery. Supports REST API and microservices infrastructure, with AIOps exposure in alert correlation and operational runbook automation.'
profile['skill_sections']=[
('Cloud and containers','Microsoft Azure, AKS, Kubernetes, Docker, Helm, ACR, Linux, private networking, Entra ID, workload identity, Key Vault'),
('Infrastructure and automation','Terraform, IaC, Azure DevOps, Azure Pipelines, GitHub Actions, Argo CD, GitOps, Ansible, Python, Bash, PowerShell, Azure CLI'),
('Reliability and operations','High availability, scalability, capacity planning, HPA, cluster autoscaler, disaster recovery readiness, backup and restore, incident troubleshooting, runbooks'),
('Monitoring and security','Azure Monitor, Log Analytics, Application Insights, Prometheus, Grafana, OpenTelemetry, ELK, RBAC, network policies, TLS, container and dependency scanning'),
('APIs and AI operations','REST APIs, FastAPI, API Management, C#/.NET and Java microservices, SQL Server, Azure SQL, Milvus, Qdrant, GPU capacity, AIOps alert correlation')]
# Reference PDF governs historical tool claims; branch profile supplies confirmed Azure operations details.
rows=[('AT&T Services Inc.','Lead Azure DevOps Engineer','08/2023 - Present',None,[
'Engineered Azure infrastructure with Terraform, provisioning AKS, ACR, private networking, Key Vault, and workload identities for isolated application environments.',
'Administered AKS namespaces, RBAC, workload identity, and resource quotas; resolved scheduling, storage, networking, image-pull, and resource-pressure incidents.',
'Planned AKS upgrades with API compatibility checks, maintenance windows, surge capacity, disruption budgets, and staged validation to protect service availability.',
'Tuned resource requests, limits, HPA, cluster autoscaler settings, and GPU node-pool capacity to balance workload scalability, performance, and cost.',
'Operated Milvus and Qdrant on AKS with persistent storage, replication, snapshots, backup procedures, and tested restores to support recovery readiness.',
'Built Azure DevOps pipelines for .NET and Java REST microservices, integrating automated tests, image scanning, Helm releases, approvals, and rollback validation.',
'Applied container hardening, least-privilege access, network policies, restricted egress, Key Vault secrets, and TLS controls across cloud-native workloads.',
'Instrumented services with Azure Monitor, Application Insights, OpenTelemetry, Prometheus, and Grafana to investigate latency, errors, resource utilization, and capacity.',
'Governed API access through Azure API Management authentication and usage limits; coordinated database migrations, deployment sequencing, and connectivity validation.',
'Developed AIOps workflows correlating Kubernetes, GPU, database, and inference alerts with deployment context and approved recovery runbooks with audit trails.'
]),('GlobalLogic','Senior DevOps Engineer','03/2023 - 08/2023',None,[
'Built reusable Azure DevOps YAML pipelines for .NET and Java services, integrating build agents, automated tests, and artifact publication.',
'Configured environment approvals and deployment verification to support controlled promotion and repeatable production releases.',
'Deployed Docker-based REST microservices to AKS with Helm and environment-specific configuration.',
'Managed Kubernetes secrets, ingress settings, and health probes to support secure connectivity and application readiness.',
'Provisioned Azure resources through Terraform modules with reviewed plans, remote state, and environment separation.',
'Maintained controlled pipeline service connections and coordinated access requirements for infrastructure deployments.',
'Coordinated SQL Server and Azure SQL schema changes with application releases, validating migration order and database permissions.',
'Investigated build, container, and Kubernetes failures and verified application and database connectivity after deployment.',
'Integrated image scanning and branch controls into delivery workflows and documented recovery procedures for failed releases.'
]),('Akamai Technologies','Senior DevOps Engineer II','12/2021 - 03/2023',None,[
'Deployed Kubernetes microservices through Argo CD GitOps workflows using version-controlled configuration.',
'Provisioned infrastructure with Terraform and pipeline-driven Infrastructure as Code deployment patterns.',
'Implemented GitHub Actions workflows for service deployments across multiple Kubernetes environments.',
'Built Docker images and maintained container packaging for application releases.',
'Packaged Kubernetes deployments with Helm charts and environment-specific values.',
'Maintained Kubernetes manifests to keep service configuration consistent across deployment environments.',
'Implemented disaster recovery controls and documented recovery requirements for cloud-hosted workloads.',
'Implemented backup controls to support infrastructure and application recovery operations.',
'Maintained observability controls to monitor workload behavior and support operational troubleshooting.'
]),('Turvo India Pvt Ltd','Senior DevOps Engineer / Senior SRE','04/2020 - 12/2021',None,[
'Designed CI/CD pipelines for application and infrastructure deployments with repeatable build and release workflows.',
'Developed internal Python FastAPI services to integrate Jenkins, Jira, and JFrog operational processes.',
'Automated release coordination and platform workflows through API-based integration with engineering tools.',
'Created custom Prometheus exporters in Python to collect application-specific metrics for SRE operations.',
'Used Prometheus and Grafana telemetry to investigate application and platform reliability issues.',
'Built Ansible roles for highly available Grafana deployments with consistent service configuration.',
'Automated Prometheus and Node Exporter deployment through reusable Ansible playbooks.',
'Authored Helm charts for repeatable Kubernetes application packaging and deployment.',
'Automated Linux and deployment operations with Python and Ansible and documented support and recovery procedures.'
]),('HCL Technologies','Senior DevOps Engineer','02/2019 - 04/2020','Mercedes-Benz Research & Development India',[
'Designed ELK monitoring architecture to collect application, infrastructure, and platform metrics.',
'Used centralized logs to investigate application behavior and support deployment troubleshooting.',
'Built Kubernetes environments with pods, Deployments, and Services for containerized applications.',
'Managed ConfigMaps to maintain environment-specific Kubernetes application configuration.',
'Packaged releases with Helm for consistent delivery across multiple namespaces.',
'Built continuous delivery pipelines for REST-based microservices using Jenkins and Maven.',
'Supported deployments across test, security validation, training, UAT, and production environments.',
'Used PowerShell scripting to automate repeatable cloud resource administration and provisioning tasks.',
'Coordinated application and deployment troubleshooting with engineering teams to verify service recovery.'
]),('Mphasis Ltd','DevOps Engineer','08/2016 - 12/2018','HPE',[
'Created Ansible deployment pipelines to promote application releases from development through production.',
'Authored Terraform modules to standardize repeatable infrastructure provisioning.',
'Executed infrastructure changes through automated CI/CD workflows.',
'Defined Git branching strategies to support release governance and controlled code promotion.',
'Documented build and deployment workflows for core infrastructure components.',
'Documented API deployment dependencies and operational handoffs for application teams.',
'Maintained database and UI release documentation to improve consistency across component deployments.',
'Developed Python and PowerShell scripts for deployment preparation, administration, and environment diagnostics.',
'Troubleshot Linux, build, configuration, and release failures with development and operations teams.'
])]
profile['experience']=[dict(company=c,title=t,date_range=dt,client=cl,impact=b) for c,t,dt,cl,b in rows]
assert profile['headline']=='Lead Azure DevOps Engineer | AI | K8s'
assert rows[0][1]=='Lead Azure DevOps Engineer'
assert all(len(x[-1])>=9 for x in rows)
Path('tmp/cloud-operations-word/profile.json').write_text(json.dumps(profile,indent=2))
d=Document(); sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.55);sec.left_margin=sec.right_margin=Inches(.65)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
styles=d.styles
for name in ['Normal','List Bullet']:
 s=styles[name];s.font.name='Calibri';s.font.size=Pt(10.5);s.paragraph_format.space_after=Pt(3);s.paragraph_format.line_spacing=1.0
styles['Title'].font.name='Calibri';styles['Title'].font.size=Pt(21);styles['Title'].font.color.rgb=RGBColor(0,0,0)
styles['Title'].paragraph_format.space_after=Pt(2)
for name in ['Heading 1','Heading 2']:
 s=styles[name];s.font.name='Calibri';s.font.size=Pt(11.5);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(4)
styles['List Bullet'].paragraph_format.left_indent=Inches(.14);styles['List Bullet'].paragraph_format.first_line_indent=Inches(-.12)
d.add_paragraph(profile['full_name'],'Title')
p=d.add_paragraph(profile['headline']);p.runs[0].bold=True;p.paragraph_format.space_after=Pt(3)
d.add_paragraph('Dallas, TX | +1 (469) 926-4201 | rajendran.scm@gmail.com')
d.add_paragraph('linkedin.com/in/rajendranelakurthi | rajendranelakurthi.github.io')
d.add_paragraph('Professional Summary','Heading 1');d.add_paragraph(profile['summary'])
d.add_paragraph('Core Skills','Heading 1')
for label,val in profile['skill_sections']:
 p=d.add_paragraph();p.add_run(label+': ').bold=True;p.add_run(val)
d.add_paragraph('Professional Experience','Heading 1')
for i,e in enumerate(profile['experience']):
 if i in (1,4):d.add_page_break()
 p=d.add_paragraph(style='Heading 2');p.add_run(e['title']+' | '+e['company'])
 p=d.add_paragraph(e['date_range']);p.paragraph_format.keep_with_next=True
 if e['client']:
  p=d.add_paragraph('Client: '+e['client']);p.paragraph_format.keep_with_next=True
 for b in e['impact']:
  p=d.add_paragraph(b,'List Bullet');p.paragraph_format.widow_control=True
# Keep credentials and education easy to parse without graphics.
d.add_paragraph('Certifications','Heading 1')
for c in profile['certifications']:d.add_paragraph(c)
d.add_paragraph('Education','Heading 1')
d.add_paragraph('Bachelor of Technology in Information Technology | 2009 - 2013')
d.add_paragraph('Prakasam Engineering College, Jawaharlal Nehru Technological University, Kakinada')
d.core_properties.title='Rajendra P N Infrastructure and Platform Engineer Resume';d.core_properties.author='Rajendra P N'
for border in list(d.styles.element.iter(qn('w:pBdr'))):
 border.getparent().remove(border)
d.save(out)
manifest={'branch':'feature/rajendrapn-azure','person':'Rajendra P N','domain':'azure-devops','level':'Tailored','source':'/Users/rajendra/Documents/1Resume/APPLIED/rajendra-pn-ai-azure-devops.pdf','profile':str(ROOT/'tmp/cloud-operations-word/profile.json'),'header':profile['headline'],'latest_title':rows[0][1],'output':str(out),'format':'docx','research':['https://learn.microsoft.com/en-us/azure/architecture/reference-architectures/containers/aks/baseline-aks'],'limitations':['No banking, core banking, event-platform or OpenShift experience added without evidence.','Historical non-Azure infrastructure generalized without recasting as Azure experience.']}
Path('tmp/cloud-operations-word/manifest.json').write_text(json.dumps(manifest,indent=2))
r=Document(out);assert not r.tables;assert len(r.inline_shapes)==0
print(out)
