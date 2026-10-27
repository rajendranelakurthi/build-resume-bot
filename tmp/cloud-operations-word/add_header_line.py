from pathlib import Path
import json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path.cwd();src=Path('/Users/rajendra/Documents/1Resume/APPLIED/docx/Rajendra_P_N_Infrastructure_Platform_Cloud_Operations.docx')
out=root/'output/docx/Rajendra_P_N_Infrastructure_Platform_Cloud_Operations_Header_Line.docx'
d=Document(src);before=[p.text for p in d.paragraphs]
assert d.paragraphs[1].text=='Lead Azure DevOps Engineer | AI | K8s'
assert any(p.text=='Lead Azure DevOps Engineer | AT&T Services Inc.' for p in d.paragraphs)
idx=next(i for i,p in enumerate(d.paragraphs) if p.text=='Professional Summary')
p=d.paragraphs[idx-1];pr=p._p.get_or_add_pPr()
for el in list(pr.findall(qn('w:pBdr'))):pr.remove(el)
bdr=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom')
for k,v in {'val':'single','sz':'6','space':'7','color':'555555'}.items():bottom.set(qn('w:'+k),v)
bdr.append(bottom);pr.append(bdr);d.save(out)
assert before==[p.text for p in Document(out).paragraphs]
m=json.loads((root/'tmp/cloud-operations-word/manifest.json').read_text());m['output']=str(out);m['formatting_update']='Horizontal divider below introductory contact block and above Professional Summary';m.pop('qa',None)
(root/'tmp/cloud-operations-word/header-line-manifest.json').write_text(json.dumps(m,indent=2))
print(out)
