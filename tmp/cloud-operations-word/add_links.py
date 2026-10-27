from pathlib import Path
import json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
root=Path.cwd();src=root/'output/docx/Rajendra_P_N_Infrastructure_Platform_Cloud_Operations_Header_Line.docx';out=src.with_name('Rajendra_P_N_Cloud_Operations_Final.docx')
d=Document(src);p=next(p for p in d.paragraphs if 'linkedin.com/in/' in p.text);p.clear()
urls=['https://www.linkedin.com/in/rajendranelakurthi','https://rajendranelakurthi.github.io']
for i,url in enumerate(urls):
 if i:p.add_run(' | ')
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');pr.append(c);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);p._p.append(h)
assert d.paragraphs[1].text=='Lead Azure DevOps Engineer | AI | K8s'
assert any(x.text=='Lead Azure DevOps Engineer | AT&T Services Inc.' for x in d.paragraphs)
d.save(out)
r=Document(out);assert all(url in [rel.target_ref for rel in r.part.rels.values() if rel.reltype==RT.HYPERLINK] for url in urls)
m=json.loads((root/'tmp/cloud-operations-word/header-line-manifest.json').read_text());m['output']=str(out);m['hyperlinks']=urls;m.pop('qa',None);(root/'tmp/cloud-operations-word/final-manifest.json').write_text(json.dumps(m,indent=2))
print(out)
