from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import json
root=Path.cwd();src=root/'output/docx/Rajendra_P_N_Cloud_Operations_Final.docx';out=src.with_name('Rajendra_P_N_Cloud_Operations_Updated.docx')
d=Document(src)
for h in d.element.iter(qn('w:hyperlink')):
 rid=h.get(qn('r:id'));target=d.part.rels[rid].target_ref
 text=''.join(t.text or '' for t in h.iter(qn('w:t')))
 if target.startswith('https://'):
  label=target.removeprefix('https://').removeprefix('www.')
  nodes=list(h.iter(qn('w:t')));nodes[0].text=label
  for n in nodes[1:]:n.text=''
assert d.paragraphs[1].text=='Lead Azure DevOps Engineer | AI | K8s'
assert any(p.text=='Lead Azure DevOps Engineer | AT&T Services Inc.' for p in d.paragraphs)
d.save(out)
m=json.loads((root/'tmp/cloud-operations-word/final-manifest.json').read_text());m['output']=str(out);m['hyperlink_display']='Domain and path only; HTTPS retained in hyperlink target';m.pop('qa',None);(root/'tmp/cloud-operations-word/updated-manifest.json').write_text(json.dumps(m,indent=2))
print(out)
