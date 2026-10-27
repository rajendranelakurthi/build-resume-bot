from docx import Document
from pathlib import Path
import json
src=Path('/Users/rajendra/Downloads/RajendraPN_DevOps_Multicloud (002).docx');d=Document(src);p=list(d.paragraphs)
def edit(i,text):
 r=p[i].runs[0];r.text=text
 for r in p[i].runs[1:]:r.text=''
changes={
7:'Builds secure, repeatable delivery platforms by combining cloud automation, repository governance, application security testing, and collaboration across development, security, and operations teams.',
10:'Lead DevOps engineer with 10+ years of experience in DevSecOps, AWS/Azure cloud infrastructure, Kubernetes, CI/CD automation, and production support using Python, Bash, PowerShell, and Groovy.',
12:'Develops CI/CD platforms with Jenkins, GitLab, and GitHub Actions, integrating automated testing, security analysis, controlled deployments, and traceable release artifacts.',
13:'Builds reusable Terraform, CloudFormation, and Ansible automation for cloud infrastructure, security, networking, and container platforms; deploys microservices through Kubernetes and Helm.',
14:'Develops Python/FastAPI integrations and Bash/Groovy automation for repository onboarding, pipeline administration, diagnostics, and enterprise DevSecOps toolchains.',
15:'Supports source control and CI/CD migrations, including Jenkins-to-GitLab workflows and GitLab-to-Enterprise GitHub repositories, coordinating tooling configuration and developer adoption.',
16:'Integrates security feedback into delivery workflows using SonarQube, Fortify, Black Duck, JFrog Xray, infrastructure policy checks, and vulnerability remediation practices.',
30:'Review/Security Tools\t\t: Fortify, SonarQube, Black Duck, JFrog Xray, Nessus, tfsec, OPA, Sentinel, Pylint; SAST and dependency vulnerability analysis',
37:'IDE and developer tooling\t: Eclipse, VSCode; AI-assisted development platform access, policy, and adoption concepts',
58:'Developed Jenkins and GitLab CI/CD pipelines for microservices, incorporating automated tests, security-analysis stages, versioned artifacts, and controlled release promotion.',
59:'Provisioned Azure Databricks workspaces across subscriptions with Terraform, standardizing environment configuration, network integration, and access controls.',
60:'Administered Azure AKS platforms, troubleshooting workloads, permissions, and networking while supporting reliable application deployment.',
61:'Maintained Azure build and release pipelines for backend services and Databricks, managing deployment inputs and environment-specific configuration.',
62:'Maintained Helm deployments for Milvus stateful services and APIs, validating configuration, readiness, and release recovery prerequisites.',
63:'Integrated continuous testing and SonarQube feedback into CI/CD workflows, working with developers to investigate failed checks and resolve release blockers.',
64:'Automated AWS/Azure infrastructure provisioning with Terraform and cloud-init, reviewing configuration changes and service dependencies before deployment.',
66:'Developed Ansible playbooks for deployment and configuration management, standardizing repeatable infrastructure changes and operational handoffs.',
73:'Administered Windows/Linux self-hosted agents and VMSS agent pools, troubleshooting dependencies, credentials, and pipeline execution failures.',
74:'Managed artifact-feed views and upstream sources, enabling consistent dependency access and versioned build inputs for application teams.',
75:'Maintained Docker build and deployment workflows, diagnosing container configuration and runtime failures during release promotion.',
76:'Administered Kubernetes workloads, Deployments, and Services, troubleshooting application health and platform connectivity.',
77:'Migrated Jenkins pipelines to GitLab CI and configured runners, validating build stages, artifact outputs, and deployment behavior during transition.',
78:'Administered GitLab repository workflows and coached developers on branching, merge reviews, access practices, and pipeline ownership.',
88:'Built GitLab CI pipelines for Node.js and Java Spring Boot applications, standardizing tests, artifacts, and environment promotion.',
89:'Automated GitLab runner provisioning on AWS EKS with Helm, providing scalable execution infrastructure for CI/CD workloads.',
91:'Migrated an on-premises ASP.NET application to AWS ECS Fargate, modernizing container delivery and supporting demand-based infrastructure scaling.',
97:'Improved image security with Nessus assessments and automated AMI refreshes through Lambda, coordinating vulnerability remediation and updated images.',
98:'Integrated GitHub Actions into application delivery, supporting consistent build, validation, and deployment workflows.',
99:'Partnered with developers on containerized application delivery, reviewing deployment prerequisites and incorporating security feedback into release readiness.',
100:'Automated AMI creation and distribution across AWS accounts with CIS and CVE checks, validating images before production adoption.',
101:'Integrated Terraform security and policy validation with tfsec GitHub Actions, OPA, and Sentinel to identify unsafe infrastructure changes before deployment.',
114:'Automated project onboarding to enterprise DevOps platforms through APIs, standardizing tool configuration and delivery workflow setup.',
115:'Developed Python/FastAPI integrations for Jenkins, Jira, and JFrog, automating cross-tool workflows and platform administration.',
117:'Supported enterprise tool ownership for CloudBees Jenkins, Stack Overflow, Discourse, and Xray, coordinating platform requirements and operational improvements.',
118:'Administered CloudBees Jenkins, JFrog Artifactory, and Xray, supporting platform availability, artifact access, and vulnerability-analysis workflows.',
122:'Developed reusable Jenkins pipelines for Docker image creation and versioned publication to private JFrog repositories, improving build consistency and traceability.',
124:'Translated platform improvements into Jira epics and stories and coordinated daily delivery discussions across teams on two continents.',
140:'Developed Python tooling to export Jenkins job configurations for plugin review, helping assess toolchain dependencies and administration requirements.',
143:'Maintained reusable Docker build images for Python, Maven, Gradle, and NPM, standardizing build dependencies across application pipelines.',
145:'Migrated source repositories from GitLab to Enterprise GitHub, validating repository contents and coordinating access, workflow changes, and developer handoffs.',
157:'Configured certificate-based authentication for MongoDB clusters, validating trusted application-to-database connectivity.',
158:'Maintained host-level TLS certificates and infrastructure access prerequisites, supporting secure production service communication.',
163:'Developed Python API automation to inspect Jenkins configurations and support repeatable CI/CD platform administration.',
164:'Managed application ACL requests and security-group rules, restricting network access to required service ports.',
167:'Developed shell automation for deployments and recurring operations, making support procedures repeatable and easier to troubleshoot.'
}
for i,t in changes.items():edit(i,t)
out=Path('output/docx/RajendraPN_DevSecOps_Multicloud_Updated.docx');d.save(out)
assert d.paragraphs[1].text==Document(src).paragraphs[1].text
assert d.paragraphs[52].text==Document(src).paragraphs[52].text
Path('tmp/devsecops-word/manifest.json').write_text(json.dumps({'source':str(src),'output':str(out.resolve()),'domain':'attachment-multicloud-devsecops','level':'Tailored','branch':'feature/rajendrapn-azure','scope':'Standalone supplied Word document; branch defaults unchanged','header':d.paragraphs[1].text,'latest_title':d.paragraphs[52].text,'ai_administration':'Concepts only; no unverified employment claim'},indent=2))
print(out.resolve())
