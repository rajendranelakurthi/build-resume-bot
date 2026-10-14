from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from copy import deepcopy
from pathlib import Path
src=Path('/Users/rajendra/Downloads/RajendraPN_AzureDevOps (1).docx')
d=Document(src);p=list(d.paragraphs)
def change(i,text):
 r=p[i].runs[0]
 r.text=text
 for other in p[i].runs[1:]:other.text=''
change(2,'Rajendra P N\nLead Azure DevOps Engineer | AI | K8s')
p[2].runs[0].font.size=Pt(14)
change(8,p[8].text.replace('Having 9 years','Having 10+ years'))
change(38,p[38].text.replace('Lead DevOps Engineer','Lead Azure DevOps Engineer'))
change(48,'Lead Azure DevOps Engineer')
change(51,p[51].text+', Azure Databricks, Azure API Management (APIM), Application Gateway, Helm, Argo CD')
change(54,'Designed and operated AKS platform infrastructure with Terraform for Milvus and GPU workloads, troubleshooting deployments, networking, resource capacity, and access controls to support reliable enterprise services.')
change(58,'Engineered repeatable Azure Databricks workspace provisioning with Terraform across isolated subscriptions, standardizing network integration, RBAC, and environment configuration for data platform teams.')
change(64,'Standardized microservice delivery using Helm, Kubernetes manifests, and GitOps patterns, integrating automated deployment checks and environment configuration into CI/CD workflows.')
new=[
'Designed, provisioned, and supported Azure API Management (APIM) gateway infrastructure with Terraform and CI/CD, maintaining API routing, backend connectivity, and environment-specific configuration.',
'Administered APIM policies for authentication, rate limiting, and request handling; investigated gateway errors and backend failures using platform diagnostics and application logs.',
'Integrated Azure Application Gateway with APIM, configuring TLS, backend health probes, and routing rules to provide secure, resilient API access.'
]
anchor=p[64]
for text in new:
 el=deepcopy(p[64]._p);anchor._p.addnext(el);q=Paragraph(el,anchor._parent)
 q.runs[0].text=text
 for r in q.runs[1:]:r.text=''
 anchor=q
out=Path('output/docx/RajendraPN_AzureDevOps_Platform_Updated.docx');d.save(out)
r=Document(out)
assert 'Lead Azure DevOps Engineer | AI | K8s' in r.paragraphs[2].text
assert r.paragraphs[48].text=='Lead Azure DevOps Engineer'
assert len(r.paragraphs)==len(p)+3
print(out.resolve())
